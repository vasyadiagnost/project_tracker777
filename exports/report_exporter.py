from pathlib import Path
import textwrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


class ReportExporter:
    @staticmethod
    def _resolve_font():
        candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/dejavu/DejaVuSans.ttf",
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/calibri.ttf",
            "C:/Windows/Fonts/tahoma.ttf",
        ]
        for font_path in candidates:
            if Path(font_path).exists():
                try:
                    pdfmetrics.registerFont(TTFont("AppReportFont", font_path))
                    return "AppReportFont"
                except Exception:
                    continue
        return "Helvetica"

    @staticmethod
    def export_txt(file_path: str, text: str) -> None:
        Path(file_path).write_text(text, encoding="utf-8")

    @staticmethod
    def export_xlsx(file_path: str, title: str, rows: list[dict]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"

        headers = ["Сигнал", "Задача", "Проект", "Этап", "DRI", "Контролёр", "Срок", "Контроль", "Статус", "Следующий шаг"]

        title_fill = PatternFill("solid", fgColor="DCE6F1")
        header_fill = PatternFill("solid", fgColor="5B9BD5")
        header_font = Font(bold=True, color="FFFFFF")
        title_font = Font(bold=True, size=14)
        thin = Side(style="thin", color="B7C9E2")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws.merge_cells("A1:J1")
        ws["A1"] = title
        ws["A1"].font = title_font
        ws["A1"].fill = title_fill
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws["A2"] = "Количество строк:"
        ws["B2"] = len(rows)

        for idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row_idx, row in enumerate(rows, start=5):
            values = [
                row.get("signal", ""),
                row.get("title", ""),
                row.get("project", ""),
                row.get("stage", ""),
                row.get("dri", ""),
                row.get("controller", ""),
                row.get("due_date", ""),
                row.get("control_date", ""),
                row.get("status", ""),
                row.get("next_step", ""),
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border

        widths = [22, 36, 24, 20, 18, 18, 14, 14, 16, 38]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:J{max(4, len(rows) + 4)}"
        wb.save(file_path)

    @staticmethod
    def export_pdf(file_path: str, title: str, lines: list[str]) -> None:
        font_name = ReportExporter._resolve_font()
        c = canvas.Canvas(file_path, pagesize=landscape(A4))
        width, height = landscape(A4)

        c.setTitle(title)

        c.setFillColor(colors.HexColor("#DCE6F1"))
        c.rect(24, height - 50, width - 48, 30, fill=1, stroke=0)
        c.setFillColor(colors.black)
        c.setFont(font_name, 14)
        c.drawString(30, height - 40, title[:140])

        y = height - 70
        c.setFont(font_name, 10)

        for line in lines:
            wrapped = textwrap.wrap(line, width=120) or [""]
            for part in wrapped:
                if y < 36:
                    c.showPage()
                    c.setFont(font_name, 10)
                    y = height - 30
                c.drawString(30, y, part)
                y -= 14

        c.save()
